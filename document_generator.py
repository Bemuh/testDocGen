import logging
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING

from case_resolver import ResolvedCase, Step


def resource_path(relpath: str | Path) -> Path:
    """Devuelve la ruta absoluta correcta dentro o fuera de PyInstaller."""
    base = Path(getattr(sys, "_MEIPASS", Path.cwd()))
    return (base / relpath).resolve()


class FontConfig:
    """ConfiguraciÇün tipogrÇ­fica."""
    def __init__(
        self,
        table_font_name,
        table_font_size,
        pasos_font_name,
        pasos_font_size,
        step_font_name,
        step_font_size,
        expected_font_name,
        expected_font_size,
        exito_font_name,
        exito_font_size,
    ):
        self.table_font_name = table_font_name
        self.table_font_size = table_font_size
        self.pasos_font_name = pasos_font_name
        self.pasos_font_size = pasos_font_size
        self.step_font_name = step_font_name
        self.step_font_size = step_font_size
        self.expected_font_name = expected_font_name
        self.expected_font_size = expected_font_size
        self.exito_font_name = exito_font_name
        self.exito_font_size = exito_font_size


class TestCaseDocumentConfig:
    """ConfiguraciÇün del documento de caso de prueba."""
    def __init__(
        self,
        header_image,
        footer_image,
        project_name,
        analyst_name,
        date,
        success_message,
        font_config,
        template_path=None,
        privacy_classification="DOCUMENTO PRIVADO",
        resultado="Éxito",
        evidencia="",
        version="001",
        case_overrides=None,
    ):
        self.header_image = header_image
        self.footer_image = footer_image
        self.project_name = project_name
        self.analyst_name = analyst_name
        self.date = date
        self.success_message = success_message
        self.font_config = font_config
        self.template_path = template_path
        self.privacy_classification = privacy_classification
        self.resultado = resultado
        self.evidencia = evidencia
        self.version = version
        self.case_overrides = case_overrides or {}
        self.user_story = ""  # se setea dinÇ­micamente


def _apply_font(run, name: str, size: int, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold


def _add_centered_image(container, image_ref, width_inches: float):
    img_path = resource_path(image_ref)

    if not img_path.exists():
        logging.warning("La imagen '%s' no se encontrÇü; se omitirÇ­.", img_path)
        return

    para = container.paragraphs[0]
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = para.add_run()
    run.add_picture(str(img_path), width=Inches(width_inches))


def _create_document_structure() -> Document:
    """Create a new document with the proper table structure for test cases.
    
    This function builds the document programmatically, exactly replicating
    the formatting from the template file.
    """
    from docx.shared import Twips
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    
    doc = Document()

    # Configure Normal style: Source Sans Pro + idioma es-ES
    normal = doc.styles['Normal']
    normal.font.name = "Source Sans Pro"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), "Source Sans Pro")
    existing_lang = rpr.find(qn('w:lang'))
    if existing_lang is not None:
        rpr.remove(existing_lang)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'es-ES')
    lang.set(qn('w:eastAsia'), 'es-ES')
    lang.set(qn('w:bidi'), 'ar-SA')
    rpr.append(lang)

    # Set page margins (1.25" sides, 1" top/bottom)
    section = doc.sections[0]
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    
    # 5-Column Grid Definitions (Total 6.0 inches)
    # Col 0: Step Number width (0.4")
    # Col 1: Remainder of Label width (0.75") -> Col 0+1 = 1.15" (Label Width)
    # Col 2: ID Value Width (1.40")
    # Col 3: Historia Label Width (1.45")
    # Col 4: Historia Value Width (2.00")
    
    col0_w = Inches(0.40)
    col1_w = Inches(0.75)
    col2_w = Inches(1.40)
    col3_w = Inches(1.45)
    col4_w = Inches(2.00)
    
    label_w = col0_w + col1_w  # 1.15"
    merged_val_w = col2_w + col3_w + col4_w # 4.85"
    step_content_w = col1_w + col2_w + col3_w + col4_w # 5.6"

    # Create header with table layout (Text | Image) for proper alignment
    header = section.header
    # Use a table for header layout: Col 1 for Text, Col 2 for Image
    htable = header.add_table(rows=1, cols=2, width=Inches(6.0))
    htable.style = None 
    htable.autofit = False 
    htable.allow_autofit = False
    
    # Set column widths (Text ~4", Image ~2")
    htable.rows[0].cells[0].width = Inches(4.0)
    htable.rows[0].cells[1].width = Inches(2.0)
    
    # Initialize text placeholder
    htable.cell(0, 0).paragraphs[0].text = "Casos De Prueba:"
    
    # Add centered title "Formato Caso de Prueba"
    title_para = doc.add_paragraph()
    title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title_run = title_para.add_run("Formato Caso de Prueba")
    title_run.bold = True
    title_run.font.name = "Source Sans Pro"
    title_run.font.size = Pt(16)
    
    # Create main table with 5 columns to support complex grid
    table = doc.add_table(rows=9, cols=5)
    table.style = 'Table Grid'
    table.autofit = False 
    table.allow_autofit = False
    
    # Helper to set cell with Source Sans Pro font
    def _set_cell_text(cell, text, bold=False, font_size=12):
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(text)
        run.bold = bold
        run.font.name = "Source Sans Pro"
        run.font.size = Pt(font_size)
    
    # Row 0: Proyecto | Value
    row0 = table.rows[0]
    # Label: Merge Col 0,1
    row0.cells[0].merge(row0.cells[1])
    _set_cell_text(row0.cells[0], "Proyecto", bold=True)
    row0.cells[0].width = label_w
    # Value: Merge Col 2,3,4
    row0.cells[2].merge(row0.cells[3]).merge(row0.cells[4])
    row0.cells[2].width = merged_val_w
    
    # Row 1: ID caso de prueba | val | Historia de usuario | val
    row1 = table.rows[1]
    # ID Label: Merge Col 0,1
    row1.cells[0].merge(row1.cells[1])
    _set_cell_text(row1.cells[0], "ID caso de prueba", bold=True)
    row1.cells[0].width = label_w
    # ID Val: Col 2
    row1.cells[2].width = col2_w
    # Hist Label: Col 3
    _set_cell_text(row1.cells[3], "Historia de usuario", bold=True)
    row1.cells[3].width = col3_w
    # Hist Val: Col 4
    row1.cells[4].width = col4_w
    
    # Row 2: Nombre del Caso de Prueba | Value
    row2 = table.rows[2]
    row2.cells[0].merge(row2.cells[1])
    _set_cell_text(row2.cells[0], "Nombre del Caso de Prueba", bold=True)
    row2.cells[0].width = label_w
    row2.cells[2].merge(row2.cells[3]).merge(row2.cells[4])
    row2.cells[2].width = merged_val_w
    
    # Row 3: Analista QA | Value
    row3 = table.rows[3]
    row3.cells[0].merge(row3.cells[1])
    _set_cell_text(row3.cells[0], "Analista QA", bold=True)
    row3.cells[0].width = label_w
    row3.cells[2].merge(row3.cells[3]).merge(row3.cells[4])
    row3.cells[2].width = merged_val_w
    
    # Row 4: Fecha | Value
    row4 = table.rows[4]
    row4.cells[0].merge(row4.cells[1])
    _set_cell_text(row4.cells[0], "Fecha", bold=True)
    row4.cells[0].width = label_w
    row4.cells[2].merge(row4.cells[3]).merge(row4.cells[4])
    row4.cells[2].width = merged_val_w
    
    # Row 5: Pasos (full width)
    row5 = table.rows[5]
    row5.cells[0].merge(row5.cells[1]).merge(row5.cells[2]).merge(row5.cells[3]).merge(row5.cells[4])
    para5 = row5.cells[0].paragraphs[0]
    para5.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run5 = para5.add_run("Pasos")
    run5.bold = True
    run5.font.name = "Source Sans Pro"
    run5.font.size = Pt(12)
    row5.cells[0].width = label_w + merged_val_w
    
    # Row 6: Template step row (step number | step content)
    row6 = table.rows[6]
    # Step Num: Col 0 (centered)
    _set_cell_text(row6.cells[0], "1", bold=True)
    row6.cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    row6.cells[0].width = col0_w
    # Content: Merge Col 1,2,3,4
    row6.cells[1].merge(row6.cells[2]).merge(row6.cells[3]).merge(row6.cells[4])
    row6.cells[1].width = step_content_w
    
    # Row 7: Resultados del caso de prueba | Value
    row7 = table.rows[7]
    row7.cells[0].merge(row7.cells[1])
    _set_cell_text(row7.cells[0], "Resultados del caso de prueba", bold=True)
    row7.cells[0].width = label_w
    row7.cells[2].merge(row7.cells[3]).merge(row7.cells[4])
    _set_cell_text(row7.cells[2], "Éxito")
    row7.cells[2].width = merged_val_w
    
    # Row 8: Evidencia | (empty)
    row8 = table.rows[8]
    row8.cells[0].merge(row8.cells[1])
    _set_cell_text(row8.cells[0], "Evidencia", bold=True)
    row8.cells[0].width = label_w
    row8.cells[2].merge(row8.cells[3]).merge(row8.cells[4])
    row8.cells[2].width = merged_val_w
    
    # Add Control De Acceso section (siempre en página nueva)

    # Control De Acceso title
    control_access_title = doc.add_paragraph()
    control_access_title.paragraph_format.page_break_before = True
    control_access_title.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run_cat = control_access_title.add_run("Control De Acceso ")
    run_cat.bold = True
    run_cat.font.name = "Source Sans Pro"
    run_cat.font.size = Pt(14)
    
    # Documento Público
    para_pub = doc.add_paragraph()
    para_pub.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run_pub = para_pub.add_run("Documento Público: ")
    run_pub.bold = True
    run_pub.font.name = "Source Sans Pro"
    run_pub.font.size = Pt(12)
    run_pub2 = para_pub.add_run("Información sin restricciones de reproducción, comunicación o transmisión completa o de cualquiera de sus partes interna o a terceros.")
    run_pub2.font.name = "Source Sans Pro"
    run_pub2.font.size = Pt(12)
    
    # Documento Privado
    para_priv = doc.add_paragraph()
    para_priv.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run_priv = para_priv.add_run("Documento Privado: ")
    run_priv.bold = True
    run_priv.font.name = "Source Sans Pro"
    run_priv.font.size = Pt(12)
    run_priv2 = para_priv.add_run("Información de uso interno. Disponible para todos los procesos internos autorizados. Esta información no puede ser conocida por terceros sin autorización formal y expresa.")
    run_priv2.font.name = "Source Sans Pro"
    run_priv2.font.size = Pt(12)
    
    # Documento Confidencial
    para_conf = doc.add_paragraph()
    para_conf.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run_conf = para_conf.add_run("Documento Confidencial: ")
    run_conf.bold = True
    run_conf.font.name = "Source Sans Pro"
    run_conf.font.size = Pt(12)
    run_conf2 = para_conf.add_run("Información privilegiada. Disponible solo para la prestación de los servicios o la operación del proceso responsable. No es permitida la reproducción, comunicación o transmisión sin la debida autorización formal y expresa del CEO, CTO.")
    run_conf2.font.name = "Source Sans Pro"
    run_conf2.font.size = Pt(12)
    
    doc.add_paragraph()  # Spacer
    
    # Access control table
    access_table = doc.add_table(rows=2, cols=2)
    access_table.style = 'Table Grid'
    _set_cell_text(access_table.rows[0].cells[0], "Acceso", bold=True)
    _set_cell_text(access_table.rows[0].cells[1], "Observaciones", bold=True)
    _set_cell_text(access_table.rows[1].cells[0], "Personal Autorizado, lideres y equipo QA")
    _set_cell_text(access_table.rows[1].cells[1], "<>")
    
    doc.add_paragraph()  # Spacer
    
    # Control De Versiones title
    control_ver_title = doc.add_paragraph()
    control_ver_title.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run_cvt = control_ver_title.add_run("Control De Versiones ")
    run_cvt.bold = True
    run_cvt.font.name = "Source Sans Pro"
    run_cvt.font.size = Pt(14)
    
    # Version table
    version_table = doc.add_table(rows=2, cols=5)
    version_table.style = 'Table Grid'
    _set_cell_text(version_table.rows[0].cells[0], "Versión", bold=True)
    _set_cell_text(version_table.rows[0].cells[1], "Fecha", bold=True)
    _set_cell_text(version_table.rows[0].cells[2], "Aprobó", bold=True)
    _set_cell_text(version_table.rows[0].cells[3], "Cargo", bold=True)
    _set_cell_text(version_table.rows[0].cells[4], "Notas", bold=True)
    _set_cell_text(version_table.rows[1].cells[0], "01")
    _set_cell_text(version_table.rows[1].cells[1], "2025/12/16")
    _set_cell_text(version_table.rows[1].cells[2], "Ana María Garzón")
    _set_cell_text(version_table.rows[1].cells[3], "QA Senior")
    _set_cell_text(version_table.rows[1].cells[4], "Creación")

    # Centrar (horizontal + vertical) todo el contenido de la tabla Control De Versiones
    for vrow in version_table.rows:
        for vcell in vrow.cells:
            vcell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for vpara in vcell.paragraphs:
                vpara.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    return doc


def _load_template(cfg: TestCaseDocumentConfig) -> Document:
    """Load the template document or create one programmatically if not found."""
    if cfg.template_path:
        template_path = resource_path(cfg.template_path)
        if template_path.exists():
            return Document(str(template_path))
        logging.warning("No se encontró el template '%s'; se creará un documento nuevo.", template_path)

    # Force programmatic creation to ensure we use the defined structure and data
    # explicitly avoiding any default .docx file on disk to remove dependency.
    logging.info("Generando estructura de documento programáticamente (sin template externo).")
    return _create_document_structure()


def _set_paragraph_text(paragraph, text: str) -> None:
    paragraph.text = text


def _set_cell_lines(cell, lines: list[str]) -> None:
    if not lines:
        cell.text = ""
        return
    cell.text = lines[0]
    for line in lines[1:]:
        cell.add_paragraph(line)


def _get_unique_cells(row):
    """Get unique cells from a row, handling merged cells.
    
    In Word/python-docx, merged cells appear multiple times in row.cells
    but share the same underlying XML element (_tc). This function
    returns only the unique cells.
    """
    seen = set()
    unique_cells = []
    for cell in row.cells:
        cell_id = id(cell._tc)  # Use the XML element ID as unique identifier
        if cell_id not in seen:
            seen.add(cell_id)
            unique_cells.append(cell)
    return unique_cells


def _clear_cell(cell) -> None:
    """Clear all content from a cell by removing paragraph text and runs."""
    for paragraph in cell.paragraphs:
        # Clear the paragraph text (this removes the first run's text)
        paragraph.clear()


def _set_cell_label(cell, text: str) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Source Sans Pro"
    run.font.size = Pt(12)


def _set_cell_value(cell, text: str) -> None:
    _clear_cell(cell)
    run = cell.paragraphs[0].add_run(text)
    run.font.name = "Source Sans Pro"
    run.font.size = Pt(12)


def _populate_header(doc: Document, resolved: ResolvedCase, cfg: TestCaseDocumentConfig) -> None:
    header = doc.sections[0].header
    
    # If we have a table (created by _create_document_structure), use it
    if header.tables:
        table = header.tables[0]
        left_cell = table.cell(0, 0)
        
        # Clear existing paragraphs in left cell
        for paragraph in left_cell.paragraphs:
            paragraph.text = ""
            
        lines = [
            "Casos De Prueba",
            f"Proyecto {resolved.project}",
            resolved.header_line,
            resolved.privacidad,
        ]
        
        # Populate text in left cell
        for idx, line in enumerate(lines):
            p = None
            if idx < len(left_cell.paragraphs):
                p = left_cell.paragraphs[idx]
                _set_paragraph_text(p, line)
            else:
                p = left_cell.add_paragraph(line)
            
            # Strict spacing controls
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        
        # Populate Header Image in right cell
        if cfg.header_image:
            img_path = resource_path(cfg.header_image)
            if img_path.exists() and len(table.rows[0].cells) > 1:
                right_cell = table.cell(0, 1)
                for p in right_cell.paragraphs: p.text = "" # Clear
                p = right_cell.paragraphs[0]
                p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
                r = p.add_run()
                # Use natural size (don't resize)
                r.add_picture(str(img_path))
        return

    # Fallback for non-table headers (legacy or template)
    for paragraph in header.paragraphs:
        paragraph.text = ""
    if header.paragraphs:
        header.paragraphs[0].text = "\n".join(
            [
                "Casos De Prueba:",
                f"Proyecto {resolved.project}",
                resolved.header_line,
                resolved.privacidad,
            ]
        )


def _populate_footer(doc: Document, cfg: TestCaseDocumentConfig) -> None:
    if not cfg.footer_image:
        return
        
    img_path = resource_path(cfg.footer_image)
    if not img_path.exists():
        return
        
    footer = doc.sections[0].footer
    # Clear existing footer content
    for p in footer.paragraphs: p.text = ""
    
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER  # Align footer image to center
    r = p.add_run()
    # Use natural size (don't resize)
    r.add_picture(str(img_path))


def _find_main_table(doc: Document):
    for table in doc.tables:
        content = " ".join(cell.text for row in table.rows for cell in row.cells)
        has_project = "Proyecto" in content or "Nombre del Proyecto" in content
        has_case = "Nombre del Caso de Prueba" in content or "CP-" in content
        if has_project and has_case and len(table.rows) >= 5:
            return table
    return doc.tables[0] if doc.tables else None


def _ensure_row_count(table, count: int) -> None:
    while len(table.rows) < count:
        table.add_row()


def _remove_row(table, row_idx: int) -> None:
    table._tbl.remove(table.rows[row_idx]._tr)


def _clone_row(row):
    return deepcopy(row._tr)


def _populate_steps(table, pasos_row_idx: int, result_row_idx: int, steps: list[Step]):
    if pasos_row_idx is None or result_row_idx is None or pasos_row_idx >= result_row_idx:
        return

    # Capture the template row (Row 6 in our structure) before removing potential placeholders
    template_row = table.rows[pasos_row_idx + 1] if pasos_row_idx + 1 < len(table.rows) else None
    
    # Remove existing placeholder rows between Pasos and Results
    for idx in range(result_row_idx - 1, pasos_row_idx, -1):
        _remove_row(table, idx)

    insert_before = table.rows[pasos_row_idx + 1] if pasos_row_idx + 1 < len(table.rows) else table.rows[-1]
    
    for step_index, step in enumerate(steps, start=1):
        if template_row is None:
            new_row = table.add_row()
        else:
            new_tr = _clone_row(template_row)
            insert_before._tr.addprevious(new_tr)
            new_row = table.rows[pasos_row_idx + step_index]

        cells = new_row.cells
        if cells:
            cells[0].text = ""
            para = cells[0].paragraphs[0]
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = para.add_run(str(step_index))
            run.bold = True
            run.font.name = "Source Sans Pro"
            run.font.size = Pt(12)

        if len(cells) > 1:
            action_text = step.action.strip()
            expected_text = step.expected.strip()
            
            cells[1].text = ""  # Clear cell
            action_paragraph = cells[1].paragraphs[0]
            
            # Action text
            action_run = action_paragraph.add_run(action_text)
            action_run.bold = True
            action_run.font.name = "Source Sans Pro"
            action_run.font.size = Pt(12)
            
            # Expected result
            if expected_text:
                # Add a line break manually to keep it in same paragraph/cell
                action_paragraph.add_run("\n")
                
                expected_run = action_paragraph.add_run(expected_text)
                expected_run.font.name = "Source Sans Pro"
                expected_run.font.size = Pt(12)


def _pick_target_table(doc: Document):
    if not doc.tables:
        return None
    for table in doc.tables:
        if not table.rows or not table.rows[0].cells:
            continue
        head = table.rows[0].cells[0].text
        if "Nombre del Proyecto" in head or "Proyecto" in head:
            return table
    return doc.tables[0]


def _populate_main_table(doc: Document, resolved: ResolvedCase) -> None:
    table = _pick_target_table(doc)
    if table is None:
        return

    pasos_row_idx = None
    result_row_idx = None
    evidencia_row_idx = None

    _ensure_row_count(table, 5)
    
    # Use unique cells to handle merged cells properly
    row0_cells = _get_unique_cells(table.rows[0])
    if len(row0_cells) >= 2:
        if "Proyecto" not in row0_cells[0].text:
            _set_cell_label(row0_cells[0], "Proyecto")
        
        # Specially format the Project value with Source Sans Pro
        _clear_cell(row0_cells[1])
        r = row0_cells[1].paragraphs[0].add_run(resolved.project)
        r.font.name = "Source Sans Pro"
        r.font.size = Pt(12)

    row1_cells = _get_unique_cells(table.rows[1])
    if len(row1_cells) >= 2:
        if "ID caso de prueba" not in row1_cells[0].text:
             _set_cell_label(row1_cells[0], "ID caso de prueba")
        
        # Format Case ID value
        _clear_cell(row1_cells[1])
        r = row1_cells[1].paragraphs[0].add_run(resolved.case_id)
        r.font.name = "Source Sans Pro"
        r.font.size = Pt(12)

    if len(row1_cells) >= 4:
        if "Historia de usuario" not in row1_cells[2].text:
             _set_cell_label(row1_cells[2], "Historia de usuario")
        
        # Format User Story value
        historia_val = getattr(resolved, 'historia_usuario', '')
        _clear_cell(row1_cells[3])
        r = row1_cells[3].paragraphs[0].add_run(historia_val)
        r.font.name = "Source Sans Pro"
        r.font.size = Pt(12)

    row2_cells = _get_unique_cells(table.rows[2])
    if len(row2_cells) >= 2:
        if "Nombre del Caso de Prueba" not in row2_cells[0].text:
            _set_cell_label(row2_cells[0], "Nombre del Caso de Prueba")
        _set_cell_value(row2_cells[1], resolved.title)

    row3_cells = _get_unique_cells(table.rows[3])
    if len(row3_cells) >= 2:
        if "Analista QA" not in row3_cells[0].text:
             _set_cell_label(row3_cells[0], "Analista QA")
        _set_cell_value(row3_cells[1], resolved.analyst)

    row4_cells = _get_unique_cells(table.rows[4])
    if len(row4_cells) >= 2:
        if "Fecha" not in row4_cells[0].text:
             _set_cell_label(row4_cells[0], "Fecha")
        _set_cell_value(row4_cells[1], resolved.date_slash)

    # First pass: only find pasos_row_idx (before step population changes indices)
    for idx, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        if not cells:
            continue
        label = cells[0]
        label_norm = " ".join(label.replace("\xa0", " ").split()).lower()
        if label_norm == "pasos":
            pasos_row_idx = idx
            break  # Found pasos, no need to continue

    # Find result_row_idx before step population (needed for _populate_steps)
    for idx, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        if not cells:
            continue
        label = cells[0]
        label_norm = " ".join(label.replace("\xa0", " ").split()).lower()
        if label_norm.startswith("resultados del caso de prueba"):
            result_row_idx = idx
            break

    # Populate steps (this changes row indices!)
    _populate_steps(table, pasos_row_idx, result_row_idx, resolved.steps)

    # After step population, re-detect result/evidencia rows with correct indices
    result_row_idx = None
    evidencia_row_idx = None
    
    for idx, row in enumerate(table.rows):
        unique_cells = _get_unique_cells(row)
        if len(unique_cells) < 2:
            continue
        cell_texts = [c.text.strip().lower() for c in unique_cells]
        
        # Detect "Resultados del caso de prueba" row (anchor on the label only)
        if result_row_idx is None:
            label_norm = " ".join(cell_texts[0].replace("\xa0", " ").split())
            if label_norm.startswith("resultados del caso de prueba"):
                result_row_idx = idx
        
        # Detect "Evidencia" row
        if evidencia_row_idx is None and result_row_idx is not None and idx > result_row_idx:
            label_norm = " ".join(cell_texts[0].replace("\xa0", " ").split())
            if label_norm == "evidencia":
                evidencia_row_idx = idx
            elif any(ext in " ".join(cell_texts) for ext in [".mp4", ".png", ".jpg", ".jpeg", ".gif", ".avi", ".mov"]):
                evidencia_row_idx = idx
            elif cell_texts[0] == "" and idx == result_row_idx + 1:
                evidencia_row_idx = idx
    
    # Create rows if not found
    if result_row_idx is None:
        _ensure_row_count(table, len(table.rows) + 1)
        result_row_idx = len(table.rows) - 1
    if evidencia_row_idx is None:
        _ensure_row_count(table, result_row_idx + 2)
        evidencia_row_idx = result_row_idx + 1

    # Now set the labels and values with correct row indices
    _ensure_row_count(table, result_row_idx + 1)
    row = table.rows[result_row_idx]
    row_cells = _get_unique_cells(row)
    if len(row_cells) >= 2:
        if "Resultados del caso de prueba" not in row_cells[0].text:
             _set_cell_label(row_cells[0], "Resultados del caso de prueba")
        _set_cell_value(row_cells[1], resolved.resultado)

    _ensure_row_count(table, evidencia_row_idx + 1)
    row = table.rows[evidencia_row_idx]
    row_cells = _get_unique_cells(row)
    if len(row_cells) >= 2:
        if "Evidencia" not in row_cells[0].text:
            _set_cell_label(row_cells[0], "Evidencia")
        _set_cell_value(row_cells[1], resolved.evidencia)


def create_test_case_document(resolved: ResolvedCase, cfg: TestCaseDocumentConfig) -> Document:
    doc = _load_template(cfg)
    _populate_header(doc, resolved, cfg)
    _populate_main_table(doc, resolved)
    _populate_footer(doc, cfg)
    return doc
