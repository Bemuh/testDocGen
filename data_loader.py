from pathlib import Path
from typing import Dict, List, Tuple, Union
from openpyxl import load_workbook

Step = Tuple[str, str]
CaseDict = Dict[str, Dict[str, List[Step]]]

def _normalize_case_id(value: Union[str, int, float]) -> str:
    if isinstance(value, (int, float)):
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text.replace(".", "", 1).isdigit():
        return text[:-2]
    return text


def load_test_cases_from_excel(path: Path) -> CaseDict:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    # ubica índices de columnas por nombre
    head = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
    id_idx  = head.index("ID")
    tit_idx = head.index("Title")
    act_idx = head.index("Step Action")
    exp_idx = head.index("Step Expected")

    cases: CaseDict = {}
    current = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        r_id, r_tit, r_act, r_exp = row[id_idx], row[tit_idx], row[act_idx], row[exp_idx]

        if r_id and r_tit:
            current = _normalize_case_id(r_id)
            cases[current] = {"title": r_tit, "steps": []}

        if current and r_act:
            cases[current]["steps"].append((str(r_act), str(r_exp or "")))

    return cases
